import streamlit as st
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

# Inicijalizacija memorije za literaturu
if "literatura" not in st.session_state:
    st.session_state.literatura = []

def kreiraj_referat(naslov, predmet, mentor, ucenici, razred, datum, podnaslovi):
    doc = Document()

    # --- 1. NASLOVNICA ---
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.add_run("Gimnazija Karlovac\n\n\n\n\n\n\n")

    run_naslov = p_header.add_run(f"{naslov.upper()}\n")
    run_naslov.bold = True
    run_naslov.font.size = Pt(16)
    p_header.add_run(f"Referat iz {predmet}\n\n\n\n\n\n\n\n\n")

    p_info = doc.add_paragraph()
    p_info.add_run(f"Mentor/ica: {mentor}\n")
    p_info.add_run(f"Učenik/ca: {ucenici}, {razred}\n\n\n\n\n\n")

    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer.add_run(f"Karlovac, {datum}.")
    
    doc.add_page_break()

    # --- 2. SADRŽAJ ---
    doc.add_heading('Sadržaj', level=1)
    popis_podnaslova = [p.strip() for p in podnaslovi.split('\n') if p.strip()]
    
    stranica = 3
    for p in popis_podnaslova:
        doc.add_paragraph(f"{p}\t{stranica}")
        stranica += 1
    doc.add_paragraph(f"Popis literature i izvora\t{stranica}")
    doc.add_page_break()

    # --- 3. POGLAVLJA ---
    for p in popis_podnaslova:
        doc.add_heading(p, level=1)
        doc.add_page_break()

    # --- 4. LITERATURA ---
    doc.add_heading("Popis literature i izvora", level=1)
    
    if st.session_state.literatura:
        # Abecedno sortiranje izvora na kraju rada
        for izvor in sorted(st.session_state.literatura):
            doc.add_paragraph(izvor)
    else:
        doc.add_paragraph("(Učenik nije unio literaturu u aplikaciji.)")

    # Spremanje u memoriju
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

# --- WEB SUČELJE ---
st.title("Generator školskog referata")

st.header("1. Osnovni podaci")
naslov = st.text_input("Naslov referata")
predmet = st.text_input("Naziv predmeta (npr. povijesti)")
mentor = st.text_input("Ime i prezime mentora (s titulom)")
ucenici = st.text_input("Ime i prezime učenika")
razred = st.text_input("Razred (npr. 3.a)")
datum = st.text_input("Datum (npr. 3. rujna 2018.)")

st.header("2. Izborne opcije")

st.subheader("Sadržaj")
podnaslovi_input = st.text_area(
    "Unesi podnaslove referata (svaki u novi red):", 
    "Uvod\nRazrada teme\nZaključak"
)

st.subheader("Literatura i izvori")
tip_izvora = st.selectbox("Odaberi vrstu izvora za unos:", ["Knjiga", "Web stranica"])

col1, col2 = st.columns(2)
with col1:
    autor = st.text_input("Autor (Prezime, Ime ili dr.)")
    godina = st.text_input("Godina izdanja")
    naslov_izvora = st.text_input("Naslov djela/članka")

with col2:
    if tip_izvora == "Knjiga":
        izdavac = st.text_input("Izdavač (npr. Školska knjiga)")
        grad = st.text_input("Mjesto izdanja (npr. Zagreb)")
    else:
        url = st.text_input("URL (poveznica)")
        datum_pristupa = st.text_input("Datum pristupa (npr. 31. kolovoza 2018.)")

if st.button("➕ Dodaj izvor u literaturu"):
    if tip_izvora == "Knjiga" and autor and naslov_izvora:
        # Formatiranje knjige
        izvor_str = f"{autor}. {godina + '. ' if godina else ''}{naslov_izvora}. {izdavac}. {grad}."
        st.session_state.literatura.append(izvor_str)
        st.success("Knjiga je dodana!")
    elif tip_izvora == "Web stranica" and naslov_izvora and url:
        # Formatiranje web stranice
        godina_str = f"{godina}. " if godina else ""
        izvor_str = f"{autor}. {godina_str}{naslov_izvora}. Preuzeto s {url} (pristupljeno {datum_pristupa})."
        st.session_state.literatura.append(izvor_str)
        st.success("Web stranica je dodana!")
    else:
        st.error("Molimo ispuni barem Autora i Naslov.")

# Prikaz dodane literature
if st.session_state.literatura:
    st.write("**Trenutno dodani izvori (bit će automatski sortirani abecedno):**")
    for lit in st.session_state.literatura:
        st.write(f"- {lit}")
    
    if st.button("🗑️ Obriši svu literaturu"):
        st.session_state.literatura = []
        st.rerun()

st.divider()

# Gumb za generiranje
if st.button("Pripremi dokument"):
    if naslov and predmet and ucenici:
        word_file = kreiraj_referat(naslov, predmet, mentor, ucenici, razred, datum, podnaslovi_input)
        ime_datoteke = f"Referat_{ucenici.replace(' ', '_')}.docx"
        
        st.success("Dokument je spreman!")
        st.download_button(
            label="⬇️ Preuzmi Referat",
            data=word_file,
            file_name=ime_datoteke,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    else:
        st.error("Ispuni barem Naslov, Predmet i Ime učenika!")
